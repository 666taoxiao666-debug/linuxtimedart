#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 TimeDART 组会汇报 PPT（16:9，学术蓝风格，基于讲稿四部分结构）"""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import numpy as np

plt.rcParams["font.sans-serif"] = ["Noto Sans CJK SC", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ── 路径 ──
ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "presentation"
OUT_DIR.mkdir(parents=True, exist_ok=True)
PPT_PATH = OUT_DIR / "TimeDART_SDWPF_组会汇报_完整版.pptx"

BEST_RUN_DIR = (
    ROOT
    / "outputs/test_results/PromptTimeDART/SDWPF/"
    "MS_il336_pl96_dm128_el2_p12_s12_lossMIXED_lr3e-05_res1_hw1.0_pw0.0_mix0.8_seed2024_id20260804_113342"
)
BLEND_DIR = ROOT / "outputs/blend_results/PromptTimeDART/SDWPF/id20260804_124707"
ASSETS = ROOT / "assets"

# ── 学术蓝配色（参考 Verifiable AI 模板）──
NAVY = RGBColor(0, 51, 102)
ROYAL = RGBColor(30, 90, 168)
SKY = RGBColor(232, 240, 248)
LIGHT_BLUE = RGBColor(210, 228, 245)
PALE_BLUE = RGBColor(245, 249, 252)
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(245, 166, 35)
GRAY = RGBColor(100, 110, 125)
LIGHT_GRAY = RGBColor(180, 190, 200)
DARK = RGBColor(40, 50, 65)
ACCENT_GREEN = RGBColor(46, 139, 87)
ACCENT_RED = RGBColor(180, 80, 80)

FONT_CN = "Noto Sans CJK SC"
FONT_EN = "Arial"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# 底部导航章节
NAV_SECTIONS = ["封面", "破局", "架构", "损失", "实验", "总结"]

PERSIST = {"mae": 222.54, "rmse": 369.49, "r2": 0.037}
BEST_MODEL = {"mae": 272.24, "rmse": 365.55, "r2": 0.057}
BEST_BLEND = {"mae": 232.47, "rmse": 343.40, "r2": 0.168}


def _set_font(run, size=14, bold=False, color=DARK, cn=True):
    run.font.name = FONT_EN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if cn:
        run.font._element.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
            FONT_CN,
        )


def add_bg(slide, color=WHITE):
    bg = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    sp = slide.shapes._spTree
    sp.remove(bg._element)
    sp.insert(2, bg._element)


def add_top_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.06)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ROYAL
    bar.line.fill.background()


def add_nav_bar(slide, active_idx):
    """底部章节导航条"""
    total_w = Inches(12.333)
    left_start = Inches(0.5)
    w = total_w / len(NAV_SECTIONS)
    for i, name in enumerate(NAV_SECTIONS):
        left = left_start + w * i
        rect = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, Inches(7.05), w, Inches(0.35)
        )
        if i == active_idx:
            rect.fill.solid()
            rect.fill.fore_color.rgb = ROYAL
            tc = WHITE
            bold = True
        else:
            rect.fill.solid()
            rect.fill.fore_color.rgb = SKY
            tc = GRAY
            bold = False
        rect.line.fill.background()
        tb = slide.shapes.add_textbox(left, Inches(7.07), w, Inches(0.32))
        tf = tb.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = name
        _set_font(r, 8, bold, tc)


def add_page_header(slide, num, title_cn, title_en=""):
    """内容页标题：编号 + 中文 + 英文副标题"""
    add_top_bar(slide)
    # 编号圆角块
    num_box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.45), Inches(0.18), Inches(0.55), Inches(0.42),
    )
    num_box.fill.solid()
    num_box.fill.fore_color.rgb = ROYAL
    num_box.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.2), Inches(0.55), Inches(0.38))
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = num
    _set_font(r, 11, True, WHITE, cn=False)

    # 标题
    tb2 = slide.shapes.add_textbox(Inches(1.1), Inches(0.15), Inches(11), Inches(0.55))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = title_cn
    _set_font(r2, 20, True, NAVY)
    if title_en:
        p3 = tf2.add_paragraph()
        r3 = p3.add_run()
        r3.text = title_en
        _set_font(r3, 10, False, GRAY, cn=False)

    # 下划线
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.45), Inches(0.72), Inches(12.4), Inches(0.025)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_BLUE
    line.line.fill.background()


def add_section_divider(slide, sec_num, title_cn, summary, nav_idx):
    """章节分隔页：大号编号 + 概述 + 导航"""
    add_bg(slide, PALE_BLUE)
    add_top_bar(slide)

    # 大号章节号
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(3), Inches(1.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = sec_num
    _set_font(r, 72, True, ROYAL, cn=False)

    # 竖线
    vline = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.4), Inches(0.06), Inches(2.2)
    )
    vline.fill.solid()
    vline.fill.fore_color.rgb = ROYAL
    vline.line.fill.background()

    # 标题
    tb2 = slide.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(0.8))
    tf2 = tb2.text_frame
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = title_cn
    _set_font(r2, 32, True, NAVY)

    # 概述框
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.6), Inches(3.2), Inches(12.1), Inches(1.6),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = LIGHT_BLUE
    box.line.width = Pt(1.5)
    tb3 = slide.shapes.add_textbox(Inches(0.85), Inches(3.4), Inches(11.6), Inches(1.3))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    r3 = p3.add_run()
    r3.text = summary
    _set_font(r3, 13, False, DARK)

    add_nav_bar(slide, nav_idx)


def add_card(slide, left, top, width, height, title, bullets, fill=SKY, border=LIGHT_BLUE):
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = border
    card.line.width = Pt(1)

    tb = slide.shapes.add_textbox(
        left + Inches(0.15), top + Inches(0.12), width - Inches(0.3), height - Inches(0.2)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    _set_font(r, 11, True, NAVY)
    for b in bullets:
        p = tf.add_paragraph()
        p.space_before = Pt(3)
        r = p.add_run()
        r.text = "▸ " + b
        _set_font(r, 9.5, False, DARK)


def add_summary_box(slide, text, top=Inches(6.35)):
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.45), top, Inches(12.4), Inches(0.62),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(230, 242, 255)
    box.line.color.rgb = ROYAL
    box.line.width = Pt(1.5)
    tb = slide.shapes.add_textbox(Inches(0.65), top + Inches(0.08), Inches(12.0), Inches(0.48))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "💡 核心结论 | Key Takeaway："
    _set_font(r, 10, True, NAVY)
    r2 = p.add_run()
    r2.text = text
    _set_font(r2, 10, False, DARK)


def add_table(slide, left, top, width, rows_data, col_widths=None, fs=9):
    n_rows = len(rows_data)
    n_cols = len(rows_data[0])
    height = Inches(0.32 * n_rows)
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    if ri == 0:
                        _set_font(r, fs, True, WHITE)
                    elif ci == 0:
                        _set_font(r, fs, True, NAVY)
                    else:
                        _set_font(r, fs, False, DARK)
            if ri == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            elif ri % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = SKY
    return table


def add_image_safe(slide, path, left, top, width=None, height=None):
    if Path(path).exists():
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
        return True
    return False


def add_flow_box(slide, left, top, w, h, text, fill=SKY, text_color=NAVY, fs=8):
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, w, h
    )
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = ROYAL
    box.line.width = Pt(1)
    tb = slide.shapes.add_textbox(left + Inches(0.05), top + Inches(0.05), w - Inches(0.1), h - Inches(0.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    _set_font(r, fs, True, text_color)


def add_arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2
    )
    conn.line.color.rgb = ROYAL
    conn.line.width = Pt(1.5)


# ── 图表生成 ──

def make_architecture_chart():
    """架构流程图"""
    fig, ax = plt.subplots(figsize=(12, 4.5), dpi=150)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 4.5)
    ax.axis("off")

    boxes = [
        (0.2, 2.8, 1.6, 0.9, "Input\n[B,336,10]\n9 exo + power", "#E8F0F8", "#003366"),
        (2.1, 2.8, 1.6, 0.9, "Channel\nIndependence\n[B×10,336,1]", "#E8F0F8", "#003366"),
        (4.0, 2.8, 1.6, 0.9, "Patch+Embed\n28 patches\n[B×10,28,128]", "#1E5AA8", "white"),
        (5.9, 2.8, 1.6, 0.9, "Causal\nTransformer×2\n(is_mask=False)", "#1E5AA8", "white"),
        (7.8, 2.8, 1.6, 0.9, "Regime\nSoft Prompt\n[B×10,128]", "#4A90D9", "white"),
        (9.7, 2.8, 1.5, 0.9, "FlattenHead\n+ Residual\n[B,96,1]", "#2E8B57", "white"),
    ]
    for x, y, w, h, txt, fc, tc in boxes:
        rect = mpatches.FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.05", fc=fc, ec="#003366", lw=1.2
        )
        ax.add_patch(rect)
        ax.text(x + w / 2, y + h / 2, txt, ha="center", va="center", fontsize=7,
                color=tc, fontweight="bold")

    for i in range(len(boxes) - 1):
        x1 = boxes[i][0] + boxes[i][2]
        x2 = boxes[i + 1][0]
        y = boxes[i][1] + boxes[i][3] / 2
        ax.annotate("", xy=(x2, y), xytext=(x1, y),
                    arrowprops=dict(arrowstyle="->", color="#003366", lw=1.5))

    # 设计目标标注
    ax.text(6.5, 1.5, "Planned: Weather CNN Prompt [B,1,128] -> QuantileHead P10/P50/P90",
            ha="center", fontsize=9, color="#F5A623", fontweight="bold",
            bbox=dict(boxstyle="round", fc="#FFF8E6", ec="#F5A623", lw=1))

    ax.text(6.5, 0.6, "Finetune excludes: Diffusion / CrossAttn Decoder / Manual Stats",
            ha="center", fontsize=8, color="#B45050", style="italic")

    ax.set_title("PromptTimeDART Finetune Forward Path", fontsize=12, fontweight="bold", pad=8)
    plt.tight_layout()
    path = OUT_DIR / "chart_architecture.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_horizon_chart():
    segments = ["Short\n1-12", "Mid-Short\n13-24", "Mid\n25-48", "Long\n49-96"]
    model_skill = [-5.8, -10.8, -9.9, 5.7]
    blend_skill = [0.0, 0.0, 1.9, 9.4]

    fig, ax = plt.subplots(figsize=(7, 3.5), dpi=150)
    x = np.arange(4)
    w = 0.35
    ax.bar(x - w / 2, model_skill, w, label="Model (Run 9)", color="#1E5AA8")
    ax.bar(x + w / 2, blend_skill, w, label="Blend (Best)", color="#2E8B57")
    ax.axhline(0, color="gray", lw=0.8)
    ax.set_xticks(x)
    ax.set_xticklabels(segments, fontsize=9)
    ax.set_ylabel("RMSE Skill vs Persistence (%)", fontsize=9)
    ax.set_title("Horizon-wise RMSE Skill vs Persistence", fontsize=11, fontweight="bold")
    ax.legend(fontsize=8)
    ax.grid(axis="y", alpha=0.3)
    plt.tight_layout()
    path = OUT_DIR / "chart_horizon_skill_v2.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_training_chart():
    runs = [
        (2, 301.22, 396.96), (4, 276.32, 373.92), (5, 273.47, 372.14),
        (9, 272.24, 365.55), (10, 273.42, 366.69),
    ]
    ids = [r[0] for r in runs]
    maes = [r[1] for r in runs]
    rmses = [r[2] for r in runs]

    fig, ax = plt.subplots(figsize=(8, 3.5), dpi=150)
    x = np.arange(len(ids))
    w = 0.35
    ax.bar(x - w / 2, maes, w, label="MAE (kW)", color="#1E5AA8", alpha=0.85)
    ax.bar(x + w / 2, rmses, w, label="RMSE (kW)", color="#4A90D9", alpha=0.75)
    ax.axhline(222.54, color="#F5A623", ls="--", lw=1.5, label="Persist MAE=222.5")
    ax.axhline(369.49, color="#E74C3C", ls=":", lw=1.5, label="Persist RMSE=369.5")
    ax.set_xticks(x)
    ax.set_xticklabels([f"Run {i}" for i in ids], fontsize=9)
    ax.set_ylabel("Error (kW)", fontsize=10)
    ax.set_title("Key Training Runs — Error Progression", fontsize=11, fontweight="bold")
    ax.legend(fontsize=7, loc="upper right")
    ax.grid(axis="y", alpha=0.3)
    plt.tight_layout()
    path = OUT_DIR / "chart_key_runs.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def make_loss_diagram():
    fig, ax = plt.subplots(figsize=(10, 3), dpi=150)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")

    boxes = [
        (0.3, 0.8, 2.8, 1.4, "L_point (active)\nMIXED: 0.8·MSE + 0.2·MAE\n+ Horizon weight w_h", "#1E5AA8", "white"),
        (3.5, 0.8, 2.8, 1.4, "L_deriv (planned)\n|dY_hat - dY| slope loss\npenalize ramp lag", "#4A90D9", "white"),
        (6.7, 0.8, 2.8, 1.4, "L_PINAW (planned)\nmax(0, PINAW - tau)\ntight interval penalty", "#2E8B57", "white"),
    ]
    for x, y, w, h, txt, fc, tc in boxes:
        rect = mpatches.FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.05", fc=fc, ec="#003366", lw=1.2
        )
        ax.add_patch(rect)
        ax.text(x + w / 2, y + h / 2, txt, ha="center", va="center", fontsize=8,
                color=tc, fontweight="bold")

    ax.text(5, 2.5, "L_total = L_quantile + lambda_d·L_deriv + lambda_p·L_PINAW  (Quantile stage)",
            ha="center", fontsize=10, fontweight="bold", color="#003366")
    ax.text(5, 0.2, "Current finetune: L = MIXED(pred, target) @ [B,96,1], residual: y_hat = P_{t-1} + dP",
            ha="center", fontsize=8, color="#666")
    plt.tight_layout()
    path = OUT_DIR / "chart_loss.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


# ── 幻灯片 ──

def slide01_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, PALE_BLUE)
    add_top_bar(slide)

    stripe = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(0.12), SLIDE_H
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ROYAL
    stripe.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12), Inches(1.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "TimeDART 风电功率预测重构"
    _set_font(r, 34, True, NAVY)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Wind Power Forecasting via PromptTimeDART on SDWPF"
    _set_font(r2, 16, False, GRAY, cn=False)

    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = "课题破局 · 架构重构 · 损失设计 · 实验验证"
    _set_font(r3, 12, False, ROYAL)

    add_card(slide, Inches(0.6), Inches(3.0), Inches(3.8), Inches(1.8),
             "数据集 | SDWPF", [
                 "输入/预测: 336 / 96 (56h→16h)",
                 "通道: 10 → 1 (MS 模式)",
                 "窗口: 144K / 21K / 51K",
             ])
    add_card(slide, Inches(4.7), Inches(3.0), Inches(3.8), Inches(1.8),
             "基线 | Persistence", [
                 f"MAE = {PERSIST['mae']} kW",
                 f"RMSE = {PERSIST['rmse']} kW",
                 f"R² = {PERSIST['r2']}",
             ], ACCENT_GREEN)
    add_card(slide, Inches(8.8), Inches(3.0), Inches(3.8), Inches(1.8),
             "最优结果 | Best", [
                 f"纯模型 Run 9: RMSE {BEST_MODEL['rmse']} kW",
                 f"Blend: RMSE {BEST_BLEND['rmse']} kW",
                 f"RMSE Skill +7.1%",
             ], GOLD)

    add_image_safe(slide, ASSETS / "2_TimeDART.png", Inches(0.6), Inches(5.1), height=Inches(1.5))
    add_nav_bar(slide, 0)


def slide02_section1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_divider(
        slide, "01", "课题破局与核心困境",
        "原生深度学习风电模型陷入「过度平滑（均值回归）」与「无效全值域置信区间」"
        "的两难困境。本部分阐述问题本质与重构动机。",
        1,
    )


def slide03_dilemma(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "1.1", "核心困境：均值回归 vs 无效区间", "Core Dilemma in Probabilistic Wind Forecasting")

    add_card(slide, Inches(0.45), Inches(0.95), Inches(5.9), Inches(2.3),
             "困境一：过度平滑 Over-smoothing", [
                 "MSE/Huber 点预测 → 条件期望 → ramp 反应迟钝",
                 "Run 2–3: R² = −0.112, MAE ≈ 301 kW",
                 "低功率高估 (MBE=+97.5 kW)，高功率低估",
                 "方向准确率仅 32.4%，低于 Persistence (35.6%)",
             ], RGBColor(255, 240, 240), RGBColor(220, 180, 180))

    add_card(slide, Inches(6.6), Inches(0.95), Inches(6.3), Inches(2.3),
             "困境二：无效置信区间 Invalid Intervals", [
                 "无约束分位数 → PICP 达标但 PINAW 爆炸",
                 "固定宽度区间 → 全值域覆盖，调度无价值",
                 "需「覆盖极端 ramp + 紧致常规段」的风险定价",
                 "→ 引出 Quantile + PINAW 复合损失（设计目标）",
             ], RGBColor(255, 248, 230), RGBColor(245, 200, 150))

    add_table(slide, Inches(0.45), Inches(3.5), Inches(12.4),
              [
                  ["症状", "量化证据", "重构对策"],
                  ["均值回归", "R²<0 (Run 1–3)", "残差预测 + MIXED 损失"],
                  ["ramp 滞后", "长期段 R²=−0.30 (Run 9)", "Regime Prompt + L_deriv (设计)"],
                  ["区间无效", "未实现概率头", "QuantileHead P10/P50/P90 + PINAW"],
              ],
              [Inches(2.5), Inches(4.5), Inches(5.4)])

    add_summary_box(slide,
        "破局核心：解耦气象外生与功率内生动力学，用 Prompt 注入上下文，"
        "残差+MIXED 修正点预测，Quantile+PINAW 服务调度风险定价。")
    add_nav_bar(slide, 1)


def slide04_section2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_divider(
        slide, "02", "模型架构与张量流转",
        "分步展示无数据泄露的前向传播链路：数据解耦 → Patch 嵌入 → "
        "Causal Transformer → Regime Prompt → FlattenHead 残差输出。",
        2,
    )


def slide05_tensor_flow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "2.1", "物理张量流转（微调路径）", "Tensor Flow without Data Leakage")

    steps = [
        ("Step 0 数据入口", "batch_x [B,336,10]\nbatch_y [B,96,10]\n→ 仅 history 进 model"),
        ("Step 1 通道解耦", "ChannelIndependence\n[B×10, 336, 1]\nMS: f_dim=−1 @ loss"),
        ("Step 2 Patch 嵌入", "Patch(12,12)→28 patches\nLinear(12→128)\n+ PosEnc"),
        ("Step 3 编码", "CausalTransformer×2\nd=128, h=8, L=2\nis_mask=False"),
    ]
    for i, (title, desc) in enumerate(steps):
        x = Inches(0.45) + Inches(3.15) * i
        add_flow_box(slide, x, Inches(0.95), Inches(2.95), Inches(1.5), f"{title}\n{desc}", SKY, NAVY, 8)

    steps2 = [
        ("Step 4 Prompt", "power hidden [B,28,128]\n→ RegimePredictor\n→ SoftPrompt [B×10,128]"),
        ("Step 5 预测头", "FlattenHead\nLinear(28×128→96)\n[B, 96, 10]"),
        ("Step 6 残差解码", "P̂ = P_{t−1} + ΔP\nzero-init head\n→ [B, 96, 1]"),
        ("设计目标", "exo→Conv1d(k=3)\n→ [B,1,128]\nQuantile P10/P50/P90"),
    ]
    for i, (title, desc) in enumerate(steps2):
        x = Inches(0.45) + Inches(3.15) * i
        fill = RGBColor(255, 248, 230) if i == 3 else WHITE
        add_flow_box(slide, x, Inches(2.7), Inches(2.95), Inches(1.5), f"{title}\n{desc}", fill, NAVY, 8)

    add_card(slide, Inches(0.45), Inches(4.5), Inches(12.4), Inches(0.9),
             "无泄露保证 | No-Leakage Guarantees", [
                 "① 微调 forward 仅接收 batch_x；batch_y 不进入 model  "
                 "② 全局时间戳 70/10/20 切分；Scaler 仅 fit 训练段  "
                 "③ 窗口不跨 TurbID / 不跨 10min 缺口",
             ], PALE_BLUE, LIGHT_BLUE)

    add_summary_box(slide,
        "当前 SDWPF: 10 维 (9 exo + power)；设计扩展至 12 维时 exo_x=x[...,:-1], target_x=x[...,-1:]。")
    add_nav_bar(slide, 2)


def slide06_architecture(prs, chart_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "2.2", "架构流程图", "Architecture Flowchart")

    add_image_safe(slide, chart_path, Inches(0.45), Inches(0.85), width=Inches(12.4))

    add_card(slide, Inches(0.45), Inches(4.0), Inches(6.0), Inches(1.2),
             "预训练路径（不在微调调用）", [
                 "Diffusion 加噪 → DenoisingPatchDecoder (CrossAttn + AdaLN)",
                 "Loss = L_recon + 0.1·CE(regime_logits, pseudo_labels)",
             ], SKY, LIGHT_BLUE)

    add_card(slide, Inches(6.7), Inches(4.0), Inches(6.15), Inches(1.2),
             "微调路径（实际调用）", [
                 "encoder → Regime Prompt → FlattenHead + Residual",
                 "Checkpoint 排除 diffusion.* / denoising_patch_decoder.*",
             ], RGBColor(230, 245, 230), ACCENT_GREEN)

    add_summary_box(slide,
        "微调阶段刻意丢弃 Diffusion/Decoder，encoder 表征可迁移；"
        "Run 4 残差头使 MAE 从 301→276 kW（−8.3%）。")
    add_nav_bar(slide, 2)


def slide07_section3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_divider(
        slide, "03", "损失函数与风险定价",
        "定量解释 MIXED 点预测损失（已实现）与一阶差分 / PINAW 约束（设计目标）"
        "的物理与工程意义。",
        3,
    )


def slide08_loss(prs, chart_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "3.1", "组合损失函数", "Combined Loss Formulation")

    add_image_safe(slide, chart_path, Inches(0.45), Inches(0.85), width=Inches(12.4))

    add_card(slide, Inches(0.45), Inches(2.5), Inches(3.9), Inches(2.0),
             "L_point — MIXED (已实现)", [
                 "L = 0.8·MSE + 0.2·|e|",
                 "MSE 惩罚大误差；MAE 减轻均值坍缩",
                 "Horizon 权重 w_h (Run 9: 均匀)",
                 "残差: ŷ = P_{t−1} + ΔP, head 零初始化",
             ])

    add_card(slide, Inches(4.6), Inches(2.5), Inches(3.9), Inches(2.0),
             "L_deriv — 一阶差分 (设计)", [
                 "L = mean|Δŷ − Δy|",
                 "直接惩罚坡度滞后",
                 "针对 ramp-up/down 高频突变",
                 "点 MSE 对此不敏感",
             ], RGBColor(255, 248, 230), GOLD)

    add_card(slide, Inches(8.75), Inches(2.5), Inches(4.1), Inches(2.0),
             "L_PINAW — 宽度约束 (设计)", [
                 "PINAW = mean(ŷ⁹⁰−ŷ¹⁰)/(y_max−y_min)",
                 "hinge: max(0, PINAW−τ)",
                 "PICP≥90% 后压缩区间",
                 "服务电网备用容量定价",
             ], RGBColor(230, 245, 230), ACCENT_GREEN)

    add_table(slide, Inches(0.45), Inches(4.75), Inches(12.4),
              [
                  ["损失项", "状态", "物理直觉", "工程意义"],
                  ["MIXED + Residual", "✅ 已实现", "修正 Persistence 增量", "Run 9 RMSE Skill +1.1%"],
                  ["L_deriv", "⏳ 设计", "跟踪 ramp 前沿", "减少爬坡滞后"],
                  ["L_PINAW", "⏳ 设计", "紧致有效区间", "可调度不确定性"],
              ],
              [Inches(2.2), Inches(1.2), Inches(3.5), Inches(5.5)], fs=8)

    add_summary_box(slide,
        "当前最优: MIXED(0.8/0.2) + 残差 + lr=3e-5；"
        "概率阶段: L = L_quantile + λ_d·L_deriv + λ_p·L_PINAW。")
    add_nav_bar(slide, 3)


def slide09_section4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_divider(
        slide, "04", "实验结果与消融分析",
        "10 次系统训练 + Persistence 融合；对比原生 DL vs 重构方案"
        "在点预测与分时段 Skill 上的改善。",
        4,
    )


def slide10_results(prs, chart_runs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "4.1", "核心实验对比", "Empirical Results")

    add_table(slide, Inches(0.45), Inches(0.85), Inches(7.5),
              [
                  ["方法", "MAE↓", "RMSE↓", "R²↑", "Skill"],
                  ["Persistence", "222.54", "369.49", "0.037", "—"],
                  ["原生 DL (Run 2–3)", "301.22", "396.96", "−0.112", "−7.4%"],
                  ["重构 Run 4 (+残差)", "276.32", "373.92", "0.014", "−1.2%"],
                  ["重构 Run 9 ★", "272.24", "365.55", "0.057", "+1.1%"],
                  ["Blend ★★", "232.47", "343.40", "0.168", "+7.1%"],
              ],
              [Inches(2.0), Inches(1.1), Inches(1.1), Inches(0.9), Inches(1.0)])

    add_image_safe(slide, chart_runs, Inches(8.2), Inches(0.85), width=Inches(4.7))

    # 改善幅度高亮
    highlights = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.45), Inches(3.2), Inches(12.4), Inches(0.85),
    )
    highlights.fill.solid()
    highlights.fill.fore_color.rgb = RGBColor(255, 248, 230)
    highlights.line.color.rgb = GOLD
    highlights.line.width = Pt(2)

    metrics = [
        ("−9.6%", "MAE 改善\nvs 原生"),
        ("−7.9%", "RMSE 改善\nvs 原生"),
        ("+0.169", "R² 提升\n绝对值"),
        ("+7.1%", "Blend\nRMSE Skill"),
    ]
    for i, (num, label) in enumerate(metrics):
        x = Inches(0.7) + Inches(3.0) * i
        circle = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, x, Inches(3.35), Inches(1.1), Inches(1.1),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = ROYAL if i < 3 else ACCENT_GREEN
        circle.line.fill.background()
        tb = slide.shapes.add_textbox(x, Inches(3.42), Inches(1.1), Inches(0.95))
        tf = tb.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = num
        _set_font(r, 14, True, WHITE, cn=False)
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = label
        _set_font(r2, 7, False, WHITE)

    add_table(slide, Inches(0.45), Inches(4.3), Inches(12.4),
              [
                  ["区间指标 (设计目标)", "原生 DL", "重构方案", "方向"],
                  ["Q-Loss (Pinball)", "—", "—", "↓ 待测"],
                  ["PICP@90%", "—", "≥90%", "↑ 待测"],
                  ["PINAW", "高 (过宽)", "低 (紧致)", "↓ 待测"],
                  ["Peak-MSE@top10%", "—", "—", "↓ 待测"],
              ],
              [Inches(3.5), Inches(2.5), Inches(2.5), Inches(3.9)], fs=8)

    add_summary_box(slide,
        "Run 9 纯模型 RMSE 365.55 kW 首次整体优于 Persistence；"
        "Blend 全局最优 RMSE 343.40 kW (Skill +7.1%)。")
    add_nav_bar(slide, 4)


def slide11_horizon_ablation(prs, chart_horizon):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "4.2", "分时段 Skill 与消融", "Horizon Skill & Ablation")

    add_image_safe(slide, chart_horizon, Inches(0.45), Inches(0.85), width=Inches(5.5))

    add_table(slide, Inches(6.2), Inches(0.85), Inches(6.65),
              [
                  ["时段", "MAE", "R²", "Skill"],
                  ["短期 1–12", "112.22", "0.809", "−5.8%"],
                  ["中短 13–24", "203.11", "0.491", "−10.8%"],
                  ["中期 25–48", "274.53", "0.076", "−9.9%"],
                  ["长期 49–96 ★", "328.39", "−0.304", "+5.7%"],
              ],
              [Inches(1.5), Inches(1.2), Inches(1.0), Inches(1.5)])

    add_table(slide, Inches(0.45), Inches(3.0), Inches(12.4),
              [
                  ["消融配置", "长期 Skill", "Ramp MAE", "结论"],
                  ["无 Prompt (基线)", "−0.7%", "—", "长期段无优势"],
                  ["+ Regime Soft Prompt", "+5.7%", "待补", "工况调制有效"],
                  ["+ Weather CNN (设计)", "Δ_skill", "Δ_peak", "预期捕获 exo 超前信号"],
              ],
              [Inches(3.0), Inches(2.0), Inches(2.0), Inches(5.4)], fs=8)

    add_image_safe(slide, BEST_RUN_DIR / "error_by_horizon.png",
                   Inches(0.45), Inches(4.0), width=Inches(5.8))
    add_image_safe(slide, BEST_RUN_DIR / "prediction_scatter.png",
                   Inches(6.5), Inches(4.0), width=Inches(5.8))

    add_summary_box(slide,
        "短期输给 Persist、长期赢 Persist — 风电预测是 Horizon 异质性问题；"
        "cutoff=24 融合是工程最优解。", top=Inches(6.55))
    add_nav_bar(slide, 4)


def slide12_section5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_divider(
        slide, "05", "总结与 Q&A 预案",
        "三条核心结论 + 组会答辩高频尖锐问题的硬核回答。",
        5,
    )


def slide13_conclusions(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "5.1", "三条金句结论", "Key Conclusions")

    conclusions = [
        ("01", "表征可迁移，解码定成败",
         "残差 + MIXED 把 R² 从 <0 拉回 >0，RMSE 首次击败 Persistence。"
         "TimeDART encoder 可迁移，损失与解码形式决定能否兑现。"),
        ("02", "Horizon 异质性 > 单一模型",
         "短期输给 Persist、长期赢 Persist。"
         "cutoff=24 融合是工程最优解，不是模型失败。"),
        ("03", "下一步：概率风险定价",
         "Quantile 头 + L_deriv + PINAW："
         "让区间在 ramp 处够窄、极端处够宽，服务调度。"),
    ]
    for i, (num, title, body) in enumerate(conclusions):
        y = Inches(0.95) + Inches(1.55) * i
        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(0.45), y, Inches(12.4), Inches(1.35),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = SKY if i % 2 == 0 else WHITE
        box.line.color.rgb = LIGHT_BLUE

        num_c = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.65), y + Inches(0.35), Inches(0.55), Inches(0.55),
        )
        num_c.fill.solid()
        num_c.fill.fore_color.rgb = ROYAL
        num_c.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(0.65), y + Inches(0.38), Inches(0.55), Inches(0.5))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = num
        _set_font(r, 14, True, WHITE, cn=False)

        tb2 = slide.shapes.add_textbox(Inches(1.4), y + Inches(0.15), Inches(11.2), Inches(1.05))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run()
        r2.text = title
        _set_font(r2, 12, True, NAVY)
        p3 = tf2.add_paragraph()
        r3 = p3.add_run()
        r3.text = body
        _set_font(r3, 10, False, DARK)

    add_nav_bar(slide, 5)


def slide14_qa(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_page_header(slide, "5.2", "Q&A 答辩预案", "Defense Preparation")

    qa_items = [
        ("Q1 未来天气泄露？",
         "微调仅 batch_x [B,336,10] 历史进 model；batch_y 不进 forward。"
         "全局时间切分 + Scaler 仅 fit 训练段；exo 是同步 SCADA 观测，非 NWP 未来预报。"),
        ("Q2 为何丢弃 Diffusion/Decoder？",
         "预训练学通用流形，微调做 96 步外推；CrossAttn Decoder 短 horizon 过参数化。"
         "encoder 可迁移 (Run 4 MAE −25 kW)；QuantileHead 替代 Diffusion 采样，O(1) 推理。"),
        ("Q3 MAE 仍差 22%，模型有用吗？",
         "看 Horizon 分层：49–96 步 Skill +5.7%，Blend 全局 +7.1%，R² 0.037→0.168。"
         "调度关心 16h 远期备用；MBE=+97.5 kW 低功率高估是 Quantile+PINAW 要解决的。"),
    ]
    for i, (q, a) in enumerate(qa_items):
        y = Inches(0.9) + Inches(1.65) * i
        add_card(slide, Inches(0.45), y, Inches(12.4), Inches(1.45),
                 q, [a], SKY if i % 2 == 0 else WHITE, LIGHT_BLUE)

    tb = slide.shapes.add_textbox(Inches(0.45), Inches(6.0), Inches(12.4), Inches(0.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = ("References: Wang et al., TimeDART, ICML 2025 · "
              "PatchTST (NeurIPS 2023) · SDWPF Wind Power Dataset")
    _set_font(r, 8, False, GRAY, cn=False)

    add_nav_bar(slide, 5)


def slide15_thanks(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, PALE_BLUE)
    add_top_bar(slide)

    tb = slide.shapes.add_textbox(Inches(0), Inches(2.5), SLIDE_W, Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "感谢聆听 · Thank You"
    _set_font(r, 36, True, NAVY)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "欢迎批评指正 · Questions & Discussion"
    _set_font(r2, 16, False, GRAY)

    add_nav_bar(slide, 5)


def main():
    print("生成辅助图表...")
    chart_arch = make_architecture_chart()
    chart_horizon = make_horizon_chart()
    chart_runs = make_training_chart()
    chart_loss = make_loss_diagram()

    print("生成 PPT...")
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide01_title(prs)
    slide02_section1(prs)
    slide03_dilemma(prs)
    slide04_section2(prs)
    slide05_tensor_flow(prs)
    slide06_architecture(prs, chart_arch)
    slide07_section3(prs)
    slide08_loss(prs, chart_loss)
    slide09_section4(prs)
    slide10_results(prs, chart_runs)
    slide11_horizon_ablation(prs, chart_horizon)
    slide12_section5(prs)
    slide13_conclusions(prs)
    slide14_qa(prs)
    slide15_thanks(prs)

    prs.save(str(PPT_PATH))
    print(f"✅ PPT 已保存: {PPT_PATH}")
    print(f"   共 {len(prs.slides)} 页 · 16:9 · 学术蓝风格")
    print(f"   辅助图表: {OUT_DIR}")


if __name__ == "__main__":
    main()
